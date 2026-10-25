"""P0 的题意路由、展示格式与保守导出回归，不调用真实模型。"""
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook, load_workbook

from table_qa.application.routing import classify
from table_qa.config import load_config
from table_qa.infrastructure.excel import export_submission
from table_qa.schemas import Question, PipelineError
from table_qa.solving.direct import solve_direct
from table_qa.solving.structured import solve_structured
from table_qa.validation.answer import canonical, format_answer


class OptimizationTest(unittest.TestCase):
    def test_field_list_repair_needs_no_arithmetic_plan(self):
        class Model:
            def ask(self, *args, **kwargs):
                return {"answer": ["名称", "金额"], "evidence": [{"page": 1, "bbox": [0, 0, 10, 10]}]}
        q = Question("1", "a", "thinking", "按顺序列出字段", answer_format="json_array", route="field_list")
        result = solve_direct(q, [{"page": 1, "width": 10, "height": 10, "image": "mock.png"}], Model(), repair={})
        self.assertEqual(result["answer"], ["名称", "金额"])
        self.assertFalse(result["computation_checked"])

    def test_calculation_route_overrides_extract_protocol(self):
        class Model:
            def ask(self, stage, payload, **kwargs):
                assert payload["question"]["solver_type"] == "thinking"
                return {"plan": {"op": "sum", "refs": ["c1"], "digits": 2}}
        q = Question("1", "a", "extract", "求和", answer_format="number", route="calculation")
        tables = [{"cells": [{"id": "c1", "text": "12.50", "sources": [{"page": 1, "bbox": [0, 0, 10, 10]}]}]}]
        result = solve_structured(q, tables, Model())
        self.assertEqual(result["answer"], "12.50")
        self.assertTrue(result["computation_checked"])

    def test_routing_semantics_and_override(self):
        cfg = load_config()
        cfg['run']['single_pass'] = False  # 保留旧路由显式compare的兼容用例
        for text, category, strategy in [
            ("该表主要使用哪种语言", "metadata", "direct"),
            ("最左侧表头", "metadata", "direct"),
            ("按表中顺序列出以下字段", "field_list", "direct"),
            ("蓝色区域的值", "visual", "direct"),
            ("求和", "calculation", "compare"),
            ("提取对应的值", "lookup", "direct")]:
            q = Question("1", "a.png", "thinking", text)
            route = classify(q, cfg)
            self.assertEqual((route["category"], route["strategy"]), (category, strategy))
            self.assertEqual(route, classify(replace(q, id="999", file_name="other.pdf"), cfg))
        cfg["run"]["strategy"] = "structured"
        self.assertEqual(classify(q, cfg)["strategy"], "structured")

    def test_hint_does_not_change_routing(self):
        q = Question("1", "a.png", "extract", "提取关键管理人员薪酬", "关联租赁情况")
        self.assertEqual(classify(q, load_config()), classify(replace(q, table_hint="任意参考"), load_config()))
        self.assertNotIn("hint_topics", load_config()["routes"])

    def test_json_count_and_generic_json(self):
        q = Question("synthetic", "a.png", "thinking", "符合条件的数量", answer_format="json")
        self.assertEqual(format_answer(q, 2), "2")
        self.assertEqual(format_answer(replace(q, question="说明关系"), {"a": 1}), '{"a":1}')
        with self.assertRaises(PipelineError):
            format_answer(replace(q, question_type="structure"), 2)

    def test_numeric_policies_and_semantic_comparison(self):
        q = Question("1", "a", "extract", "值", answer_format="number")
        self.assertEqual(format_answer(q, "１，２５０．００"), "1250.00")
        self.assertEqual(format_answer(q, "18.00%"), "18.00%")
        self.assertEqual(format_answer(q, "18%", {"percent": "fraction"}), "0.18")
        self.assertEqual(format_answer(q, "18%", {"percent": "number"}), "18")
        self.assertEqual(format_answer(q, "(12.50)"), "(12.50)")
        self.assertEqual(format_answer(q, "(12.50)", {"parentheses_negative": "convert"}), "-12.50")
        self.assertEqual(canonical(q, "18%"), canonical(q, ".18"))
        self.assertNotEqual(canonical(q, "18%"), canonical(q, "18"))
        self.assertEqual(canonical(q, "12.50"), canonical(q, "12.5"))
        for value in [True, "18%解释", "(坏值)", "(-12)"]:
            with self.assertRaises(PipelineError):
                format_answer(q, value)

    def test_brackets_are_not_conflated(self):
        q = Question("1", "a", "extract", "值", answer_format="string")
        self.assertEqual(format_answer(q, "【Ａ】 ［Ｂ］"), "【Ａ】 ［Ｂ］")
        self.assertEqual(format_answer(q, "【Ａ】 ［Ｂ］", {"normalize_text_width": True}), "【A】 [B]")

    def test_uncertain_export_requires_opt_in_and_evidence(self):
        q = Question("1", "a", "extract", "值", answer_format="number")
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = Workbook(); book.active.append(["id", "answer"])
            book.save(root / "template.xlsx"); book.close()
            record = {"status": "uncertain", "answer": "12.50", "evidence": [{"page": 1}],
                      "selected_candidate": {"answer": "12.50"}}
            args = ([q], {"1": record}, root / "template.xlsx", root / "out.xlsx")
            self.assertEqual(export_submission(*args)["blank_count"], 1)
            report = export_submission(*args, submit_uncertain=True)
            self.assertEqual(report["uncertain_submitted_count"], 1)
            self.assertEqual(report["unresolved_count"], 1)
            result = load_workbook(root / "out.xlsx")
            self.assertEqual(result.active["B2"].value, "12.50"); result.close()
            record["evidence"] = []
            self.assertEqual(export_submission(*args, submit_uncertain=True)["blank_count"], 1)
            report = export_submission(*args, submit_uncertain=True, submit_unlocalized=True)
            self.assertEqual(report["unlocalized_submitted_count"], 1)
            self.assertEqual(report["unresolved_count"], 1)
            record.update(status="success", answer="错误单位")
            self.assertEqual(export_submission(*args)["format_failed_ids"], ["1"])

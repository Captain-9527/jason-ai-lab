"""原文标记、文本保真和导出审计的合成数据回归。"""
import copy
import json
import tempfile
import unittest
from dataclasses import replace
from pathlib import Path

from openpyxl import Workbook, load_workbook

from table_qa.schemas import Question, PipelineError
from table_qa.validation.answer import canonical, format_answer, validate_candidate_answer
from table_qa.infrastructure.excel import export_submission
from table_qa.solving.direct import solve_direct
from table_qa.solving.structured import solve_structured


class LiteralAnswerTest(unittest.TestCase):
    def setUp(self):
        self.q = Question("synthetic", "sample.png", "extract", "提取目标单元格", answer_format="number", route="lookup")

    def test_markers_preserved_but_not_zero_or_calculation(self):
        for value in ("--", "-", "—", "–", "N/A", "n/a", "NA"):
            self.assertEqual(validate_candidate_answer(self.q, value), value)
            self.assertEqual(canonical(self.q, value), canonical(self.q, value))
            self.assertNotEqual(canonical(self.q, value), canonical(self.q, 0))
            for q in (replace(self.q, route="calculation"), replace(self.q, route="filter"),
                      replace(self.q, question_type="thinking")):
                with self.assertRaises(PipelineError):
                    validate_candidate_answer(q, value)
        for value in (None, "", "未知", "--解释", "NaN", "Infinity", True):
            with self.assertRaises(PipelineError):
                validate_candidate_answer(self.q, value)

    def test_text_and_array_preserve_width(self):
        text = "范围～（Ａ）【Ｂ】"
        self.assertEqual(format_answer(replace(self.q, answer_format="string"), text), text)
        self.assertEqual(json.loads(format_answer(replace(self.q, answer_format="json_array"), [text])), [text])

    def test_direct_and_structured_keep_literal(self):
        class Model:
            def ask(self, stage, *args, **kwargs):
                return ({"answer": "--", "evidence": [{"page": 1, "bbox": [0, 0, 100, 100]}]}
                        if stage == "direct" else {"operation": "extract", "values": [{"ref": "c", "text": "--"}]})
        pages = [{"page": 1, "width": 100, "height": 100, "image": "unused"}]
        self.assertEqual(solve_direct(self.q, pages, Model())["answer"], "--")
        tables = [{"table_id": "t", "row_count": 1, "col_count": 1, "cells": [
            {"id": "c", "text": "--", "row": 0, "col": 0, "rowspan": 1, "colspan": 1,
             "sources": [{"page": 1, "bbox": [0, 0, 10, 10]}]}]}]
        self.assertEqual(solve_structured(self.q, tables, Model())["answer"], "--")

    def test_export_audit_does_not_mutate_results(self):
        questions = [self.q, replace(self.q, id="formatted"), replace(self.q, id="missing")]
        records = {self.q.id: {"status": "success", "answer": "--"},
                   "formatted": {"status": "success", "answer": "1,250.00"}}
        before = copy.deepcopy(records)
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            book = Workbook(); book.active.append(["id", "answer"])
            book.save(root / "template.xlsx"); book.close()
            report = export_submission(questions, records, root / "template.xlsx", root / "out.xlsx", fill_missing=True)
            self.assertEqual(report["literal_marker_ids"], [self.q.id])
            self.assertEqual(report["placeholder_ids"], ["missing"])
            audit = json.loads((root / report["answers_file"]).read_text())["items"]
            self.assertEqual(audit[0]["submitted_answer"], "--")
            self.assertFalse(audit[0]["placeholder"])
            self.assertEqual(audit[1]["raw_answer"], "1,250.00")
            self.assertEqual(audit[1]["submitted_answer"], "1250.00")
            self.assertTrue(audit[1]["text_changed"])
            self.assertTrue(audit[2]["placeholder"])
            book = load_workbook(root / "out.xlsx")
            self.assertEqual(book.active['B2'].value, '--'); book.close()
        self.assertEqual(records, before)

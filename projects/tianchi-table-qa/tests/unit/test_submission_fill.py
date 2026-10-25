"""非空提交只影响导出，不将占位值冒充识别结果。"""
import copy
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook
from table_qa.schemas import Question
from table_qa.infrastructure.excel import export_submission
from table_qa.validation.submission import nonempty_answer


class SubmissionFillTest(unittest.TestCase):
    def test_all_contracts_have_explicit_placeholders(self):
        for kind in ("string", "number", "json", "json_array", "table"):
            q = Question(kind, "synthetic.png", "structure" if kind == "table" else "extract", "测试", answer_format="json" if kind == "table" else kind)
            text, detail = nonempty_answer(q, {})
            self.assertTrue(text.strip())
            self.assertTrue(detail["placeholder"])

    def test_selected_answer_then_candidates_then_baseline(self):
        q = Question("1", "synthetic.png", "extract", "测试", answer_format="number")
        record = {"answer": "invalid", "candidates": {"direct": {"answer": "12.50"}}}
        before = copy.deepcopy(record)
        text, detail = nonempty_answer(q, record, baseline={"answer": "99"})
        self.assertEqual(text, "12.50")
        self.assertEqual(detail["source"], "candidate:direct")
        self.assertEqual(record, before)
        self.assertEqual(nonempty_answer(q, {}, baseline={"answer": 0})[0], "0")
        self.assertEqual(nonempty_answer(q, {"answer": "7", **{k:v for k,v in record.items() if k != "answer"}})[0], "7")

    def test_review_preferred_and_invalid_candidates(self):
        q = Question("1", "synthetic.png", "extract", "测试", answer_format="number")
        record = {"review": {"preferred": "structured"}, "candidates": {
            "direct": {"answer": 1}, "structured": {"answer": 2}}}
        self.assertEqual(nonempty_answer(q, record)[0], "2")
        record["candidates"]["structured"]["answer"] = "bad"
        self.assertEqual(nonempty_answer(q, record)[0], "1")

    def test_export_roundtrip_and_input_unchanged(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            book = Workbook(); book.active.append(["id", "answer"])
            book.save(root / "template.xlsx"); book.close()
            questions = [Question(str(i), "s.png", "extract", "测试", answer_format="string") for i in range(4)]
            results = {"0": {"status": "success", "answer": "=1+1"},
                       "1": {"status": "uncertain", "answer": "候选", "evidence": []},
                       "2": {"status": "failed", "answer": None}}
            before = copy.deepcopy(results)
            report = export_submission(questions, results, root / "template.xlsx", root / "submission.xlsx",
                                       fill_missing=True, baseline_results={"2": {"answer": "基线"}})
            self.assertEqual(results, before)
            self.assertEqual(report["blank_count"], 0)
            self.assertEqual(report["placeholder_ids"], ["3"])
            audit = json.loads((root / "submission_fillings.json").read_text())
            self.assertEqual([v["source"] for v in audit["items"]], ["result", "baseline", "placeholder"])
            book = load_workbook(root / "submission.xlsx")
            self.assertEqual([book.active.cell(i, 2).value for i in range(2, 6)], ["=1+1", "候选", "基线", "无法确定"])
            self.assertEqual(book.active.cell(2, 2).data_type, "s")
            book.close()

    def test_overlong_candidate_falls_back(self):
        q = Question("1", "s.png", "extract", "测试", answer_format="string")
        text, detail = nonempty_answer(q, {"answer": "x" * 32768}, baseline={"answer": "合法"})
        self.assertEqual(text, "合法")
        self.assertTrue(detail["rejected"])

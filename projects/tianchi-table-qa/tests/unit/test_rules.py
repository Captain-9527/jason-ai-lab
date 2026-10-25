"""确定性逻辑测试，不访问模型或比赛文档内容。"""
import copy
import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from table_qa.config import load_config
from table_qa.infrastructure.excel import export_submission, read_questions
from table_qa.infrastructure.model_client import parse_json, prompt_data
from table_qa.infrastructure.storage import atomic_json, read_json
from table_qa.infrastructure.tracing import Trace
from table_qa.processing.long_table import merge_tables
from table_qa.schemas import PipelineError, Question
from table_qa.solving.calculator import calculate, execute_plan
from table_qa.validation.answer import format_answer
from table_qa.validation.input import check_inputs
from table_qa.validation.table import bbox, validate_table


def table(values):
    return {"row_count": len(values), "col_count": len(values[0]), "cells": [
        {"text": v, "row": r, "col": c, "rowspan": 1, "colspan": 1}
        for r, row in enumerate(values) for c, v in enumerate(row)]}


class RulesTest(unittest.TestCase):
    def test_structure_overlap(self):
        t = table([["a", "b"]])
        t["cells"][0]["colspan"] = 2
        with self.assertRaises(PipelineError):
            validate_table(t)

    def test_structure_out_of_bounds(self):
        t = table([["a"]])
        t["cells"][0]["rowspan"] = 2
        with self.assertRaises(PipelineError):
            validate_table(t)

    def test_unknown_distinct_from_blank(self):
        t = table([[None, ""]])
        self.assertIsNone(validate_table(t)["cells"][0]["text"])
        with self.assertRaises(PipelineError):
            format_answer(Question("1", "x", "structure", "结构"), t)

    def test_partial_structure_preserves_absolute_positions(self):
        t = table([["a", "b"], ["c", "d"]])
        t["cells"] = [t["cells"][-1]]
        result = json.loads(format_answer(Question("1", "x", "structure", "局部"), t))
        self.assertEqual((result["row_count"], result["col_count"]), (2, 2))
        self.assertEqual(result["cells"][0]["row"], 1)

    def test_strict_numeric_and_commas(self):
        q = Question("1", "x", "extract", "数字", answer_format="number")
        self.assertEqual(format_answer(q, "1,250.00"), "1250.00")
        for text in ["答案是1250", "12,50", "NaN", "Infinity"]:
            with self.assertRaises(PipelineError):
                format_answer(q, text)

    def test_array_values_not_objects(self):
        q = Question("1", "x", "extract", "数组", answer_format="json_array")
        self.assertEqual(json.loads(format_answer(q, [1, "", "名称"])), [1, "", "名称"])
        self.assertEqual(json.loads(format_answer(q, ["1,250", "001", "部门1,250"])), ["1250", "001", "部门1,250"])
        for values in [[{"value": 1}], [None], [[1]]]:
            with self.assertRaises(PipelineError):
                format_answer(q, values)

    def test_calculator_decimal_and_percent(self):
        self.assertEqual(calculate({"op": "sum"}, ["0.1", "0.2"]), "0.3")
        self.assertEqual(calculate({"op": "sum"}, ["(1,000)", "20"]), "-980")
        self.assertEqual(calculate({"op": "growth_rate", "scale": "100", "digits": 1, "suffix": "%"}, ["120", "100"]), "20.0%")
        self.assertEqual(calculate({"op": "sum"}, ["10%", "20%"]), "0.3")

    def test_zero_denominator_and_unknown_op(self):
        for plan, values in [({"op": "ratio"}, ["1", "0"]), ({"op": "eval"}, ["1"]), ({"op": "mean"}, [])]:
            with self.assertRaises(PipelineError):
                calculate(plan, values)

    def test_plan_rejects_unknown_ref(self):
        with self.assertRaises(PipelineError):
            execute_plan({"op": "sum", "refs": ["invented"]}, {})

    def test_sort_count_and_argmax(self):
        self.assertEqual(calculate({"op": "sort", "descending": True}, ["2", "10"]), ["10", "2"])
        self.assertEqual(calculate({"op": "count"}, []), "0")
        self.assertEqual(calculate({"op": "argmax"}, ["2", "10"], ["a", "b"]), "b")

    def test_bbox_strict(self):
        self.assertEqual(bbox([0, 1, 9, 10], 10, 10), [0, 1, 9, 10])
        for b in [[-1, 0, 9, 9], [0, 0, 11, 9], [0, 0, 0, 9], [0, 0, float("nan"), 1]]:
            with self.assertRaises(PipelineError):
                bbox(b, 10, 10)

    def test_long_table_header_and_overlap(self):
        left, right = table([["H"], ["A"], ["B"]]), table([["H"], ["B"], ["C"]])
        result = merge_tables(left, right, {"same_table": True, "column_order": [0],
                                           "drop_header_rows": 1, "overlap_rows": 1})
        self.assertEqual(result["row_count"], 4)
        self.assertEqual([c["text"] for c in result["cells"]], ["H", "A", "B", "C"])
        self.assertEqual(right["row_count"], 3)

    def test_long_table_rejects_false_duplicate(self):
        with self.assertRaises(PipelineError):
            merge_tables(table([["A"]]), table([["B"], ["C"]]),
                         {"same_table": True, "column_order": [0], "overlap_rows": 1})

    def test_long_table_continued_row(self):
        result = merge_tables(table([["长"]]), table([["记录"], ["下行"]]),
                              {"same_table": True, "column_order": [0], "continued_row": True})
        self.assertEqual(result["cells"][0]["text"], "长\n记录")
        self.assertEqual(result["row_count"], 2)

    def test_long_table_rejects_bad_mapping(self):
        with self.assertRaises(PipelineError):
            merge_tables(table([["a", "b"]]), table([["c", "d"]]),
                         {"same_table": True, "column_order": [0, 0]})

    def test_json_parser_strict(self):
        self.assertEqual(parse_json('```json\n{"a":1}\n```'), {"a": 1})
        for raw in ['{"a":NaN}', '解释 {"a":1}', '[]']:
            with self.assertRaises(PipelineError):
                parse_json(raw)

    def test_prompt_omits_runtime_paths(self):
        self.assertEqual(prompt_data({"sources": [{"image": "/tmp/a", "page": 1}]}), {"sources": [{"page": 1}]})

    def test_input_corrections_opt_in_and_path_escape(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / "058.pdf").touch()
            cfg = {"paths": {"files": tmp}, "run": {"corrections_enabled": False},
                   "corrections": {"file_aliases": {"58.pdf": "058.pdf"}, "question_types": {"1": "extract"}}}
            q = Question("1", "58.pdf", "错误类型", "问", answer_format="number")
            self.assertFalse(check_inputs([q], cfg)[1]["ok"])
            cfg["run"]["corrections_enabled"] = True
            self.assertTrue(check_inputs([q], cfg)[1]["ok"])
            self.assertFalse(check_inputs([Question("2", "../outside.pdf", "extract", "问")], cfg)[1]["ok"])

    def test_trace_redaction_and_stage_duration(self):
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {"DASHSCOPE_API_KEY": "example-test-secret"}):
            path = Path(tmp) / "trace.jsonl"
            trace = Trace(path, payloads=True)
            with trace.scope(question_id="1"), trace.stage("demo", {"api_key": "example-test-secret"}) as out:
                out["value"] = "example-test-secret"
            text = path.read_text()
            self.assertNotIn("example-test-secret", text)
            self.assertIn("duration_seconds", text)
            self.assertEqual(json.loads(text.splitlines()[-1])["question_id"], "1")

    def test_excel_roundtrip_and_formula_literal(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            w = Workbook(); w.active.append(["id", "answer"]); w.save(root / "template.xlsx")
            qs = [Question("001", "a.png", "extract", "问", answer_format="string"),
                  Question("002", "a.png", "extract", "问", answer_format="string")]
            result = export_submission(qs, {"001": {"status": "success", "answer": "=1+1"}}, root / "template.xlsx", root / "out.xlsx")
            self.assertEqual(result["unresolved_ids"], ["002"])
            book = load_workbook(root / "out.xlsx")
            self.assertEqual(book.active["A2"].value, "001")
            self.assertEqual(book.active["B2"].data_type, "s")
            self.assertEqual(book.active["B2"].value, "=1+1")
            book.close()

    def test_config_no_api_key(self):
        cfg = load_config()
        self.assertNotIn("api_key", cfg["model"])


if __name__ == "__main__":
    unittest.main()

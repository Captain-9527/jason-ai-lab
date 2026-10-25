"""自建图片/Excel 和固定模拟响应验证全链路，绝不联网答比赛题。"""
import copy
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import numpy as np
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw

from table_qa.application.runner import execute, export_run, compare_runs
from table_qa.application.pipeline import answer_question
from table_qa.config import load_config
from table_qa.infrastructure.storage import read_results, read_json
from table_qa.infrastructure.tracing import Trace
from table_qa.processing.document import preprocess, prepare_document, crop_image
from table_qa.processing.table import recognize_tables
from table_qa.schemas import Question


class FakeModel:
    calls = []

    def __init__(self, cfg, trace, run, used_calls=0):
        self.trace = trace

    def reset_question_budget(self):
        pass

    def ask(self, stage, payload, images=None, vision=True, fresh=False):
        self.calls.append((stage, fresh))
        self.trace.emit("mock_response", stage=stage)
        if stage == "locate":
            return {"tables": [{"bbox": [0, 0, 200, 100], "header_bbox": None}]}
        if stage == "recognize":
            return {"row_count": 2, "col_count": 2, "cells": [
                {"row": 0, "col": 0, "rowspan": 1, "colspan": 1, "text": "A"},
                {"row": 0, "col": 1, "rowspan": 1, "colspan": 1, "text": "10"},
                {"row": 1, "col": 0, "rowspan": 1, "colspan": 1, "text": "B"},
                {"row": 1, "col": 1, "rowspan": 1, "colspan": 1, "text": "20"}]}
        if stage == "structured":
            kind = payload["question"]["question_type"]
            if kind == "extract":
                return {"values": [{"ref": "t0:r0:c1", "text": "10"}]}
            if kind == "thinking":
                return {"plan": {"op": "sum", "refs": ["t0:r0:c1", "t0:r1:c1"]}}
            return {"table_id": "t0", "cell_ids": ["t0:r1:c1"]}
        if stage == "review":
            result = {"verdict": "pass", "confidence": .95, "risks": [], "regions": []}
            if payload["question"]["question_type"] == "structure":
                answer = next(iter(payload["candidates"].values()))["answer"]
                result["full_table_check"] = {"scope": "full_table", "row_count": answer["row_count"],
                                              "col_count": answer["col_count"], "page": 1,
                                              "basis": "合成图完整表格两行两列，按整表检查"}
            return result
        if stage in {"direct", "repair"}:
            return {"answer": "10", "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
        raise AssertionError(stage)


class IntegrationTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        files = self.root / "files"; files.mkdir()
        Image.new("RGB", (200, 100), "white").save(files / "sample.png")
        w = Workbook(); w.active.append(["id", "file_name", "question_type", "question", "answer_format"])
        for i, kind, question, fmt in [("1", "extract", "提取A的数值", "number"),
                                        ("2", "thinking", "数值合计", "number"),
                                        ("3", "structure", "末格结构", "json")]:
            w.active.append([i, "sample.png", kind, question, fmt])
        w.save(self.root / "questions.xlsx")
        t = Workbook(); t.active.append(["id", "answer"]); t.save(self.root / "template.xlsx")
        self.cfg = load_config()
        # 这些用例保留历史多路协议的兼容回归；新默认流程单独测试。
        self.cfg['run'].update(single_pass=False, review=True)
        self.cfg["model"]["bbox_format"] = "pixel"  # 原模拟响应明确使用上传图像素
        self.cfg["paths"] = {"questions": str(self.root / "questions.xlsx"), "files": str(files),
                             "template": str(self.root / "template.xlsx"), "cache": str(self.root / "cache"),
                             "outputs": str(self.root / "outputs")}
        # 原集成样例验证固定结构路线；自动分流另设用例。
        self.cfg["run"].update({"allow_local_processing": True, "workers": 1,
                                 "strategy": "structured", "review_policy": "always"})
        self.cfg["processing"]["preprocess"] = False
        FakeModel.calls = []

    def tearDown(self):
        self.tmp.cleanup()

    def test_conflict_preferred_candidate_skips_repair(self):
        class Preferred(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "direct":
                    return {"answer": "11", "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
                if stage == "review":
                    return {"verdict": "pass", "confidence": .95, "preferred": "structured", "risks": []}
                if stage == "repair":
                    raise AssertionError("已选中合法候选，不应再次修复")
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"]["strategy"] = "compare"
        summary = execute(self.cfg, ids=["1"], model_factory=Preferred)
        row = read_results(Path(self.cfg["paths"]["outputs"]) / summary["run_id"])["1"]
        self.assertEqual(row["answer"], "10")
        self.assertEqual(row["selection_reason"], "review_preferred")

    def test_advisory_preserved_and_preferred_used_when_answers_agree(self):
        class Advisory(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "review":
                    return {"verdict": "pass", "confidence": .9, "preferred": "structured", "risks": [
                        {"severity": "advisory", "code": "evidence_granularity", "message": "框较粗",
                         "rationale": "已通过原图逐格核对"}]}
                if stage == "repair":
                    raise AssertionError("提示不应触发修复")
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"]["strategy"] = "compare"
        result = execute(self.cfg, ids=["1"], model_factory=Advisory)
        row = read_results(Path(self.cfg["paths"]["outputs"]) / result["run_id"])["1"]
        self.assertEqual(row["status"], "success")
        self.assertEqual(row["selected_candidate"]["operation"], "extract")
        self.assertEqual(len(row["review_history"]), 1)
        self.assertIn("复核提示: 框较粗", row["warnings"])

    def test_retry_repair_precedes_initial_direct(self):
        class Retry(FakeModel):
            reviews = 0
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "review":
                    self.reviews += 1
                    return {"verdict": "retry" if self.reviews == 1 else "unknown", "confidence": .4, "risks": ["需复查"]}
                if stage == "repair":
                    return {"answer": "12", "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"]["strategy"] = "direct"
        summary = execute(self.cfg, ids=["1"], model_factory=Retry)
        row = read_results(Path(self.cfg["paths"]["outputs"]) / summary["run_id"])["1"]
        self.assertEqual(row["status"], "uncertain")
        self.assertEqual(row["answer"], "12")
        self.assertEqual(row["selection_reason"], "repair_after_retry")

    def test_unlocalized_answer_stays_uncertain_even_review_passes(self):
        class Unlocalized(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "direct":
                    return {"answer": "10", "evidence": [{"page": 1, "bbox": [0, 0, 9000, 9000]}]}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"]["strategy"] = "direct"
        summary = execute(self.cfg, ids=["1"], model_factory=Unlocalized)
        self.assertEqual(summary["states"], {"uncertain": 1})
        self.assertEqual(summary["unlocalized_submitted_count"], 1)
        self.assertEqual(summary["coordinate_warning_count"], 1)
        self.cfg["run"]["submit_unlocalized"] = False
        self.assertEqual(export_run(self.cfg, summary["run_id"])["unlocalized_submitted_count"], 1)

    def test_normalized_locate_uses_page_size(self):
        class Normalized(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "locate":
                    return {"tables": [{"bbox": [0, 0, 999, 999]}]}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["model"]["bbox_format"] = "norm999"
        pages = [{"page": 1, "width": 200, "height": 100, "image": str(self.root / "files/sample.png")}]
        tables = recognize_tables(pages, Normalized(self.cfg, Trace(self.root / "trace"), self.root), self.cfg, self.root)
        self.assertEqual(tables[0]["segments"][0]["bbox"], [0, 0, 200, 100])

    def test_invalid_locate_falls_back_to_full_page(self):
        class Invalid(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "locate":
                    return {"tables": [{"bbox": [0, 0, 9000, 9000]}]}
                return super().ask(stage, payload, images, vision, fresh)
        pages = [{"page": 1, "width": 200, "height": 100, "image": str(self.root / "files/sample.png")}]
        tables = recognize_tables(pages, Invalid(self.cfg, Trace(self.root / "trace"), self.root), self.cfg, self.root)
        self.assertEqual(tables[0]["segments"][0]["bbox"], [0, 0, 200, 100])

    def test_review_regions_and_candidate_fallback_are_converted_once(self):
        for use_review_region in (True, False):
            class Regions(FakeModel):
                reviews = 0
                def ask(self, stage, payload, images=None, vision=True, fresh=False):
                    if stage in {"direct", "repair"}:
                        return {"answer": "10", "evidence": [{"page": 1, "bbox": [250, 250, 750, 750]}]}
                    if stage == "review":
                        self.reviews += 1
                        return {"verdict": "retry" if self.reviews == 1 else "pass", "confidence": .95,
                                "risks": [], "regions": [{"page": 1, "bbox": [250, 250, 750, 750]}] if use_review_region else []}
                    return super().ask(stage, payload, images, vision, fresh)
            self.cfg["model"]["bbox_format"] = "norm1000"
            self.cfg["run"]["strategy"] = "direct"
            with patch("table_qa.application.pipeline.crop_image", wraps=crop_image) as crop:
                result = execute(self.cfg, ids=["1"], model_factory=Regions)
                self.assertEqual(result["states"], {"success": 1})
                self.assertEqual(crop.call_args.args[1], [34, 9, 166, 91])

    def test_normalized_blocks_keep_crop_offsets(self):
        class Blocks(FakeModel):
            count = 0
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "locate":
                    return {"tables": [{"bbox": [250, 0, 750, 1000], "header_bbox": [250, 0, 750, 100]}]}
                if stage == "blocks":
                    return {"blocks": [[0, 0, 1000, 600], [0, 400, 1000, 1000]]}
                if stage == "recognize":
                    self.count += 1
                    values = ["H", "10"] if self.count == 1 else ["10", "20"]
                    return {"row_count": 2, "col_count": 1, "cells": [
                        {"row": i, "col": 0, "rowspan": 1, "colspan": 1, "text": value} for i, value in enumerate(values)]}
                if stage == "boundary":
                    return {"same_table": True, "column_order": [0], "overlap_rows": 1,
                            "drop_header_rows": 0, "continued_row": False, "warnings": []}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["model"]["bbox_format"] = "norm1000"
        self.cfg["processing"]["block_height"] = 70
        pages = [{"page": 1, "width": 200, "height": 100, "image": str(self.root / "files/sample.png")}]
        tables = recognize_tables(pages, Blocks(self.cfg, Trace(self.root / "trace"), self.root), self.cfg, self.root)
        self.assertEqual(tables[0]["row_count"], 3)
        self.assertEqual(tables[0]["segments"][0]["bbox"], [50, 0, 150, 60])
        self.assertEqual(tables[0]["segments"][1]["bbox"], [50, 40, 150, 100])

    def test_auto_low_risk_skips_recognition_and_review(self):
        self.cfg["run"].update(strategy="auto", review_policy="on_risk")
        summary = execute(self.cfg, ids=["1"], model_factory=FakeModel)
        self.assertEqual(FakeModel.calls, [("direct", False)])
        row = read_results(Path(self.cfg["paths"]["outputs"]) / summary["run_id"])["1"]
        self.assertEqual(row["validation_level"], "format_evidence")
        self.assertEqual(row["status"], "success")

    def test_auto_calculation_keeps_review(self):
        class Agree(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "direct":
                    self.calls.append((stage, fresh))
                    return {"answer": "30.00", "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"].update(strategy="auto", review_policy="on_risk")
        summary = execute(self.cfg, ids=["2"], model_factory=Agree)
        self.assertIn(("review", False), FakeModel.calls)
        self.assertEqual(summary["states"], {"success": 1})

    def test_end_to_end_and_reexport(self):
        summary = execute(self.cfg, model_factory=FakeModel)
        self.assertEqual(summary["states"], {"success": 3})
        self.assertEqual(sum(s == "recognize" for s, _ in FakeModel.calls), 1)
        folder = Path(self.cfg["paths"]["outputs"]) / summary["run_id"]
        book = load_workbook(folder / "submission.xlsx")
        self.assertEqual(book.active["B2"].value, "10")
        self.assertEqual(book.active["B3"].value, "30")
        book.close()
        before = len(FakeModel.calls)
        self.assertEqual(export_run(self.cfg, summary["run_id"])["rows"], 3)
        self.assertEqual(len(FakeModel.calls), before)

    def test_resume_skips_success_and_rerun_inherits(self):
        first = execute(self.cfg, model_factory=FakeModel)
        before = len(FakeModel.calls)
        execute(self.cfg, "resume", base=first["run_id"], model_factory=FakeModel)
        self.assertEqual(len(FakeModel.calls), before)
        second = execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=FakeModel)
        data = read_results(Path(self.cfg["paths"]["outputs"]) / second["run_id"])
        self.assertEqual(data["2"]["inherited_from"], first["run_id"])
        self.assertIsNone(data["1"]["inherited_from"])
        self.assertIn(("structured", True), FakeModel.calls)
        self.assertEqual(compare_runs(self.cfg, first["run_id"], second["run_id"])["changed_count"], 0)

    def test_resume_rejects_config_change(self):
        first = execute(self.cfg, ids=["1"], model_factory=FakeModel)
        self.cfg["prompts"]["structured"] += " changed"
        with self.assertRaisesRegex(ValueError, "已变化"):
            execute(self.cfg, "resume", base=first["run_id"], model_factory=FakeModel)

    def test_rerun_rejects_source_change(self):
        first = execute(self.cfg, ids=["1"], model_factory=FakeModel)
        Image.new("RGB", (200, 100), "black").save(Path(self.cfg["paths"]["files"]) / "sample.png")
        with self.assertRaisesRegex(ValueError, "源文件已变化"):
            execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=FakeModel)

    def test_partial_run_exports_all_rows_and_reports_missing(self):
        summary = execute(self.cfg, ids=["1"], model_factory=FakeModel)
        self.assertEqual(summary["unresolved_count"], 2)
        self.assertEqual(summary["rows"], 3)

    def test_resume_cannot_silently_expand_selection(self):
        first = execute(self.cfg, ids=["1"], model_factory=FakeModel)
        with self.assertRaisesRegex(ValueError, "不能扩展"):
            execute(self.cfg, "resume", ids=["2"], base=first["run_id"], model_factory=FakeModel)

    def test_long_table_blocks_are_merged_with_verified_overlap(self):
        class Blocks(FakeModel):
            count = 0
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "blocks":
                    return {"blocks": [[0, 0, 200, 60], [0, 40, 200, 100]]}
                if stage == "recognize":
                    self.count += 1
                    values = ["H", "10"] if self.count == 1 else ["10", "20"]
                    return {"row_count": 2, "col_count": 1, "cells": [
                        {"row": i, "col": 0, "rowspan": 1, "colspan": 1, "text": v}
                        for i, v in enumerate(values)]}
                if stage == "boundary":
                    return {"same_table": True, "column_order": [0], "overlap_rows": 1,
                            "drop_header_rows": 0, "continued_row": False, "warnings": []}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["processing"]["block_height"] = 70
        trace = Trace(self.root / "blocks.jsonl")
        pages = prepare_document(Path(self.cfg["paths"]["files"]) / "sample.png", self.cfg, trace)
        tables = recognize_tables(pages, Blocks(self.cfg, trace, self.root), self.cfg, self.root / "work")
        self.assertEqual(tables[0]["row_count"], 3)
        self.assertEqual([c["text"] for c in tables[0]["cells"]], ["H", "10", "20"])

    def test_interruption_preserves_prior_success_and_all_ids(self):
        from table_qa.schemas import BudgetExceeded
        class Limited(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "structured" and payload["question"]["question_type"] == "thinking":
                    raise BudgetExceeded("模拟预算耗尽")
                return super().ask(stage, payload, images, vision, fresh)
        with self.assertRaisesRegex(RuntimeError, "已保存"):
            execute(self.cfg, model_factory=Limited)
        folder = next(Path(self.cfg["paths"]["outputs"]).iterdir())
        self.assertEqual(read_results(folder)["1"]["status"], "success")
        self.assertEqual(read_json(folder / "report.json")["rows"], 3)

    def test_comparison_conflict_repaired_and_rechecked(self):
        class Disagree(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "direct":
                    return {"answer": "11", "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
                return super().ask(stage, payload, images, vision, fresh)
        self.cfg["run"]["strategy"] = "compare"
        summary = execute(self.cfg, ids=["1"], model_factory=Disagree)
        data = read_results(Path(self.cfg["paths"]["outputs"]) / summary["run_id"])
        self.assertEqual(data["1"]["route"], "repair")
        self.assertEqual(data["1"]["answer"], "10")

    def test_low_confidence_does_not_become_success(self):
        class Unsure(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "review":
                    return {"verdict": "unknown", "confidence": .2, "risks": ["看不清"], "regions": []}
                return super().ask(stage, payload, images, vision, fresh)
        result = execute(self.cfg, ids=["1"], model_factory=Unsure)
        self.assertEqual(result["states"], {"uncertain": 1})
        self.assertEqual(result["unresolved_count"], 3)
        self.assertEqual(result["uncertain_submitted_count"], 1)
        self.cfg["run"]["submit_uncertain"] = False
        self.assertEqual(export_run(self.cfg, result["run_id"])["uncertain_submitted_count"], 1)

    def test_zero_count_has_document_evidence(self):
        class ZeroCount(FakeModel):
            def ask(self, stage, payload, images=None, vision=True, fresh=False):
                if stage == "structured":
                    return {"plan": {"op": "count", "refs": []}}
                return super().ask(stage, payload, images, vision, fresh)
        result = execute(self.cfg, ids=["2"], model_factory=ZeroCount)
        row = read_results(Path(self.cfg["paths"]["outputs"]) / result["run_id"])["2"]
        self.assertEqual(row["answer"], "0")
        self.assertEqual(row["status"], "success")
        self.assertTrue(row["evidence"])

    def test_pdf_rendering_no_model(self):
        import pymupdf
        path = self.root / "sample.pdf"
        doc = pymupdf.open(); doc.new_page(width=100, height=100); doc.save(path); doc.close()
        pages = prepare_document(path, self.cfg, Trace(self.root / "pdf-trace.jsonl"))
        self.assertEqual(len(pages), 1)
        self.assertEqual(pages[0]["width"], 250)

    def test_deskew_direction_and_coordinate_inverse(self):
        image = Image.new("RGB", (600, 400), "white")
        draw = ImageDraw.Draw(image)
        for y in range(60, 350, 50):
            draw.line([(30, y), (570, y + 40)], fill="black", width=2)
        opts = copy.deepcopy(self.cfg["processing"])
        opts["min_side"] = 100
        result, inverse, actions = preprocess(image, opts)
        angle = next(a["deskew_degrees"] for a in actions if "deskew_degrees" in a)
        self.assertGreater(angle, 0)
        forward = np.linalg.inv(np.array(inverse))
        a, b = forward @ [30, 60, 1], forward @ [570, 100, 1]
        self.assertLess(abs(a[1] - b[1]), 5)


if __name__ == "__main__":
    unittest.main()

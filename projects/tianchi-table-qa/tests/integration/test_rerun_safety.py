"""自建数据验证失败回退、方向缓存及尺寸冲突修复。"""
import hashlib
import unittest
from pathlib import Path

import pymupdf
from PIL import Image

from tests.integration import test_pipeline as fixtures
from table_qa.application.runner import execute, export_run
from table_qa.application.pipeline import answer_question
from table_qa.infrastructure.storage import atomic_json, digest, read_json, read_results
from table_qa.infrastructure.tracing import Trace
from table_qa.processing.document import prepare_document
from table_qa.schemas import Question, PipelineError, BudgetExceeded


class RerunSafetyTest(unittest.TestCase):
    setUp=fixtures.IntegrationTest.setUp
    tearDown=fixtures.IntegrationTest.tearDown

    def test_document_migration_requires_all_related_questions_and_no_old_fallback(self):
        first = execute(self.cfg, model_factory=fixtures.FakeModel)
        origin = Path(self.cfg["paths"]["outputs"]) / first["run_id"]
        original = (origin / "results.jsonl").read_bytes()
        Image.new("RGB", (200, 100), "black").save(self.root / "files/sample.png")
        with self.assertRaisesRegex(ValueError, "源文件已变化"):
            execute(self.cfg, "rerun", ids=["1", "2", "3"], base=first["run_id"], model_factory=fixtures.FakeModel)
        self.cfg["run"]["accept_document_changes"] = True
        with self.assertRaisesRegex(ValueError, "源文件已变化"):
            execute(self.cfg, "resume", ids=["1", "2", "3"], base=first["run_id"], model_factory=fixtures.FakeModel)
        with self.assertRaisesRegex(ValueError, "全部关联题"):
            execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=fixtures.FakeModel)
        class Failure(fixtures.FakeModel):
            def ask(self, *args, **kwargs):
                raise PipelineError("模拟识别失败")
        second = execute(self.cfg, "rerun", ids=["1", "2", "3"], base=first["run_id"], model_factory=Failure)
        folder = Path(self.cfg["paths"]["outputs"]) / second["run_id"]
        self.assertEqual(second["input_migration"]["changed_documents"], ["sample.png"])
        self.assertEqual(second["input_migration"]["invalidated_ids"], ["1", "2", "3"])
        self.assertEqual(second["baseline_fallback_ids"], [])
        self.assertTrue(all(r["answer"] is None for r in read_results(folder).values()))
        self.assertEqual(export_run(self.cfg, second["run_id"])["placeholder_count"], 3)
        self.assertEqual((origin / "results.jsonl").read_bytes(), original)

    def test_document_migration_cannot_accept_question_changes(self):
        from openpyxl import load_workbook
        first = execute(self.cfg, ids=["1"], model_factory=fixtures.FakeModel)
        book = load_workbook(self.cfg["paths"]["questions"])
        book.active.cell(2, 4, "修改后的问题")
        book.save(self.cfg["paths"]["questions"])
        book.close()
        self.cfg["run"]["accept_document_changes"] = True
        with self.assertRaisesRegex(ValueError, "题目清单"):
            execute(self.cfg, "rerun", ids=["1", "2", "3"], base=first["run_id"], model_factory=fixtures.FakeModel)

    def test_recognition_artifact_reuse_skips_cloud_stages(self):
        first = execute(self.cfg, ids=["1"], model_factory=fixtures.FakeModel)
        self.cfg["run"]["reuse_recognition"] = True
        self.cfg["model"]["vision"] = "qwen-new-answer-model"
        class Cached(fixtures.FakeModel):
            def ask(self, stage, *args, **kwargs):
                if stage in {"locate", "recognize", "boundary", "blocks"}:
                    raise AssertionError("识别产物应离线复用")
                return super().ask(stage, *args, **kwargs)
        second = execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=Cached)
        self.assertEqual(second["recognition_reused_documents"], 1)
        third = execute(self.cfg, "rerun", ids=["1"], base=second["run_id"], model_factory=Cached)
        self.assertEqual(third["recognition_reused_documents"], 1)

    def test_recognition_refresh_and_changed_images_disable_artifact_reuse(self):
        first = execute(self.cfg, ids=["1"], model_factory=fixtures.FakeModel)
        self.cfg["run"].update(reuse_recognition=True, refresh_stages=["recognize"])
        fixtures.FakeModel.calls = []
        second = execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=fixtures.FakeModel)
        self.assertEqual(second["recognition_reused_documents"], 0)
        self.assertIn(("recognize", False), fixtures.FakeModel.calls)
        self.cfg["run"]["refresh_stages"] = []
        self.cfg["processing"].update(preprocess=True, min_side=500)
        fixtures.FakeModel.calls = []
        third = execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=fixtures.FakeModel)
        self.assertEqual(third["recognition_reused_documents"], 0)
        self.assertIn(("recognize", False), fixtures.FakeModel.calls)

    def test_incomplete_artifacts_are_not_reused(self):
        first = execute(self.cfg, ids=["1"], model_factory=fixtures.FakeModel)
        path = (Path(self.cfg["paths"]["outputs"]) / first["run_id"] / "artifacts" /
                digest("sample.png")[:16] / "tables/structures.json")
        tables = read_json(path)
        tables[0]["incomplete"] = True
        atomic_json(path, tables)
        self.cfg["run"]["reuse_recognition"] = True
        fixtures.FakeModel.calls = []
        second = execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=fixtures.FakeModel)
        self.assertEqual(second["recognition_reused_documents"], 0)
        self.assertIn(("recognize", False), fixtures.FakeModel.calls)

    def test_block_height_override_preserves_page_cache(self):
        trace = Trace(self.root / "trace")
        source = self.root / "files/sample.png"
        before = prepare_document(source, self.cfg, trace)
        self.cfg["run"]["recognition_block_height"] = 40
        self.assertEqual(prepare_document(source, self.cfg, trace), before)

    def test_reuse_answer_cache_changes_rerun_fresh_flag(self):
        first = execute(self.cfg, ids=["1"], model_factory=fixtures.FakeModel)
        self.cfg["run"]["reuse_answer_cache"] = True
        fixtures.FakeModel.calls = []
        execute(self.cfg, "rerun", ids=["1"], base=first["run_id"], model_factory=fixtures.FakeModel)
        self.assertIn(("structured", False), fixtures.FakeModel.calls)

    def test_failed_rerun_preserves_source_and_reexports_candidate(self):
        first=execute(self.cfg,ids=["1"],model_factory=fixtures.FakeModel)
        origin=Path(self.cfg["paths"]["outputs"])/first["run_id"]
        original=(origin/"results.jsonl").read_bytes()
        class Failure(fixtures.FakeModel):
            def ask(self,*args,**kwargs): raise PipelineError("模拟超时")
        second=execute(self.cfg,"rerun",ids=["1"],base=first["run_id"],model_factory=Failure)
        row=read_results(Path(self.cfg["paths"]["outputs"])/second["run_id"])["1"]
        self.assertEqual(row["answer"],"10")
        self.assertEqual(row["status"],"uncertain")
        self.assertEqual(row["selection_reason"],"baseline_fallback")
        self.assertEqual(second["baseline_fallback_ids"],["1"])
        self.assertEqual(second["unlocalized_submitted_count"],1)
        self.assertEqual((origin/"results.jsonl").read_bytes(),original)

    def test_budget_interruption_keeps_baseline_without_masking_error(self):
        first=execute(self.cfg,ids=["1"],model_factory=fixtures.FakeModel)
        class Limited(fixtures.FakeModel):
            def ask(self,*args,**kwargs): raise BudgetExceeded("预算停止")
        with self.assertRaisesRegex(RuntimeError,"已保存"):
            execute(self.cfg,"rerun",ids=["1"],base=first["run_id"],model_factory=Limited)
        new=next(p for p in Path(self.cfg["paths"]["outputs"]).iterdir() if p.name!=first["run_id"])
        self.assertEqual(read_results(new)["1"]["selection_reason"],"baseline_fallback")

    def test_rotated_pdf_cache_and_original_preserved(self):
        source=self.root/"rotated.pdf"
        with pymupdf.open() as doc:
            p=doc.new_page(width=400,height=600)
            p.insert_text((100,550),"Synthetic orientation text for a full table "*2,rotate=90,fontsize=8)
            doc.save(source)
        original=hashlib.sha256(source.read_bytes()).hexdigest()
        trace=Trace(self.root/"trace")
        self.cfg["processing"].update(pdf_auto_orient=True,preprocess=False)
        pages=prepare_document(source,self.cfg,trace)
        self.assertEqual(pages[0]["orientation"]["clockwise_degrees"],90)
        with Image.open(pages[0]["original"]) as im:
            self.assertEqual((pages[0]["width"],pages[0]["height"]),(im.height,im.width))
        self.assertEqual(prepare_document(source,self.cfg,trace),pages)
        self.cfg["processing"]["pdf_auto_orient"]=False
        disabled=prepare_document(source,self.cfg,trace)
        self.assertNotEqual(disabled[0]["image"],pages[0]["image"])
        self.assertEqual(disabled[0]["orientation"]["clockwise_degrees"],0)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(),original)

    def test_repair_cannot_silently_shrink_agreed_full_table_dimensions(self):
        cell = {"id": "h", "text": "H", "row": 0, "col": 0, "rowspan": 1, "colspan": 1,
                "sources": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
        answer = {"row_count": 17, "col_count": 1, "cells": [cell]}
        class Shrinking(fixtures.FakeModel):
            def ask(self, stage, payload, *args, **kwargs):
                if stage == "structured":
                    return {"table_id": "t", "cell_ids": ["h"]}
                if stage in {"direct", "repair"}:
                    return {"answer": {**answer, "row_count": 17 if stage == "direct" else 1},
                            "evidence": [{"page": 1, "bbox": [0, 0, 200, 100]}]}
                return {"verdict": "retry", "confidence": .4, "risks": ["只应输出表头"]}
        self.cfg["run"]["strategy"] = "compare"
        trace = Trace(self.root / "trace")
        q = Question("s", "sample.png", "structure", "恢复表头", answer_format="json")
        pages = [{"page": 1, "width": 200, "height": 100, "image": str(self.root / "files/sample.png")}]
        row = answer_question(q, pages, [{**answer, "table_id": "t"}],
                              Shrinking(self.cfg, trace, self.root), self.cfg, trace, self.root)
        self.assertEqual(row["answer"]["row_count"], 17)
        self.assertNotIn("repair_1", row["candidates"])
        self.assertEqual(row["status"], "uncertain")

    def test_dimensions_conflict_repairs_instead_of_selecting_header_height(self):
        calls=[]
        cell={"text":"H","row":0,"col":0,"rowspan":1,"colspan":1,"id":"h",
              "sources":[{"page":1,"bbox":[0,0,200,100]}]}
        answer={"row_count":2,"col_count":1,"cells":[cell]}
        class Structure(fixtures.FakeModel):
            def ask(self,stage,payload,*args,**kwargs):
                calls.append(stage)
                if stage=="structured":return {"table_id":"t","cell_ids":["h"]}
                if stage in {"direct","repair"}:
                    return {"answer":{**answer,"row_count":1 if stage=="direct" else 2},
                            "evidence":[{"page":1,"bbox":[0,0,200,100]}]}
                if stage != "review": raise AssertionError(stage)
                return {"verdict":"pass","confidence":.99,"risks":[],
                        "full_table_check":{"scope":"full_table","row_count":2,"col_count":1,"page":1,"basis":"整表两行，输出表头一行"}}
        trace=Trace(self.root/"trace")
        q=Question("s","sample.png","structure","恢复表头",answer_format="json")
        pages=[{"page":1,"width":200,"height":100,"image":str(self.root/"files/sample.png")}]
        self.cfg["run"].update(strategy="compare",review=False)
        row=answer_question(q,pages,[{**answer,"table_id":"t"}],Structure(self.cfg,trace,self.root),self.cfg,trace,self.root)
        self.assertEqual(row["status"],"success")
        self.assertEqual(row["answer"]["row_count"],2)
        self.assertEqual(calls,["direct","structured","repair","review"])
        self.assertTrue(row["review_history"][0]["program_check"])

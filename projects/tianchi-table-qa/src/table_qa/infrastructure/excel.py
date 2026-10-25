"""只读原始题目；在模板副本中写入答案，保存后重新检查。"""
import json
import os
import tempfile
from pathlib import Path

from openpyxl import load_workbook

from ..schemas import Question, PipelineError
from ..validation.answer import format_answer, is_literal_marker
from ..validation.submission import nonempty_answer
from .storage import atomic_json


def read_questions(path):
    book = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = book.active.iter_rows(values_only=True)
        headers = list(next(rows))
        required = {"id", "file_name", "question_type", "question"}
        if not required.issubset(headers) or len(headers) != len(set(headers)):
            raise ValueError("题目表缺少必填列或包含重复表头")
        result = []
        for row in rows:
            if not any(v is not None and v != "" for v in row):
                continue
            data = {k: "" if v is None else str(v) for k, v in zip(headers, row)}
            result.append(Question(**{k: data.get(k, "") for k in Question.__dataclass_fields__ if k != "route"}))
        return result
    finally:
        book.close()


def export_submission(questions, results, template, target, submit_uncertain=False, answer_policy=None, submit_unlocalized=False,
                      fill_missing=False, baseline_results=None):
    """ID 保持文本、答案为文字值而非公式；部分运行仍导出全部题目并标记缺失数量。"""
    ids = [q.id for q in questions]
    if len(ids) != len(set(ids)) or not all(ids):
        raise ValueError("无法导出重复或空 ID")
    if set(results) - set(ids):
        raise ValueError("结果包含本题目集之外的 ID")
    book = load_workbook(template)
    sheet = book.active
    if [sheet.cell(1, c).value for c in (1, 2)] != ["id", "answer"]:
        raise ValueError("模板前两列必须为 id、answer")
    if len(book.worksheets) != 1 or sheet.max_column != 2:
        raise ValueError("提交模板必须只有一个工作表、两列")
    if sheet.max_row > 1:
        sheet.delete_rows(2, sheet.max_row - 1)
    expected = []
    unresolved = []
    uncertain_submitted = []
    format_failed = []
    unlocalized_submitted = []
    fillings = []
    answer_audit = []
    for index, q in enumerate(questions, 2):
        record = results.get(q.id)
        answer = ""
        raw_answer = (record or {}).get("answer")
        source, placeholder = "result", False
        if not record or record["status"] != "success":
            unresolved.append(q.id)
        eligible = record and (record["status"] == "success" or
                   (submit_uncertain and record["status"] == "uncertain" and
                    record.get("selected_candidate") and (record.get("evidence") or submit_unlocalized)))
        if eligible and record.get("answer") is not None:
            try:
                answer = format_answer(q, record["answer"], answer_policy)
            except (PipelineError, ValueError, TypeError, KeyError):
                format_failed.append(q.id)
            if answer and record["status"] == "uncertain":
                uncertain_submitted.append(q.id)
                if not record.get("evidence"):
                    unlocalized_submitted.append(q.id)
        if fill_missing and (not answer.strip() or len(answer) > 32767):
            answer, detail = nonempty_answer(q, record, answer_policy, (baseline_results or {}).get(q.id))
            raw_answer = detail.get("raw_answer")
            source, placeholder = detail["source"], detail["placeholder"]
            fillings.append({"id": q.id, "answer": answer, "original_status": (record or {}).get("status", "missing"), **detail})
            if record and record.get("status") == "uncertain" and not detail["placeholder"]:
                if q.id not in uncertain_submitted:
                    uncertain_submitted.append(q.id)
                if not record.get("evidence") and q.id not in unlocalized_submitted:
                    unlocalized_submitted.append(q.id)
        if len(answer) > 32767:
            raise ValueError(f"题目 {q.id} 答案超过 Excel 单元格字符上限，拒绝截断")
        for col, value in ((1, q.id), (2, answer)):
            cell = sheet.cell(index, col, value)
            cell.data_type = "s"
        expected.append((q.id, answer))
        answer_audit.append({"id": q.id, "raw_answer": raw_answer, "submitted_answer": answer,
                             "source": source, "placeholder": placeholder,
                             "literal_marker": is_literal_marker(q, raw_answer),
                             "text_changed": isinstance(raw_answer, str) and raw_answer != answer})
    target = Path(target)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx", dir=target.parent)
    os.close(fd)
    try:
        book.save(tmp)
        check = load_workbook(tmp, read_only=True, data_only=False)
        try:
            actual = [(str(a), "" if b is None else str(b)) for a, b in check.active.iter_rows(min_row=2, max_col=2, values_only=True)]
            if actual != expected:
                raise ValueError("导出后答案或 ID 复检不一致")
            if fill_missing and any(not a.strip() for _, a in actual):
                raise ValueError("非空导出复检失败")
            if any(c.data_type == "f" for row in check.active for c in row):
                raise ValueError("提交文件中出现公式")
        finally:
            check.close()
        os.replace(tmp, target)
    finally:
        book.close()
        if os.path.exists(tmp):
            os.unlink(tmp)
    audit = target.with_name(target.stem + "_fillings.json")
    atomic_json(audit, {"fill_missing": fill_missing, "items": fillings})
    answers_file = target.with_name(target.stem + "_answers.json")
    atomic_json(answers_file, {"answer_policy": answer_policy or {}, "items": answer_audit})
    return {"rows": len(ids), "unresolved_count": len(unresolved), "unresolved_ids": unresolved,
            "answers_file": answers_file.name,
            "literal_marker_ids": [v["id"] for v in answer_audit if v["literal_marker"]],
            "fill_missing": fill_missing, "filled_count": len(fillings),
            "placeholder_ids": [v["id"] for v in fillings if v["placeholder"]],
            "placeholder_count": sum(v["placeholder"] for v in fillings), "fillings_file": audit.name,
            "uncertain_submitted_count": len(uncertain_submitted), "uncertain_submitted_ids": uncertain_submitted,
            "submitted_count": sum(bool(a) for _, a in expected),
            "unlocalized_submitted_count": len(unlocalized_submitted), "unlocalized_submitted_ids": unlocalized_submitted,
            "blank_count": sum(not a for _, a in expected), "format_failed_ids": format_failed}
